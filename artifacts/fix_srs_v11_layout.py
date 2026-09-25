from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.oxml.ns import qn


DOCX = Path(__file__).resolve().parent / 'Aether_SRS_2026-09.docx'


def set_font(run, size=9.3, bold=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_after(anchor, text, style=None, after=5):
    new = OxmlElement('w:p')
    anchor._p.addnext(new)
    p = type(anchor)(new, anchor._parent)
    if style:
        p.style = style
    set_font(p.add_run(text), 9 if style else 9.3)
    p.paragraph_format.space_after = Pt(after)
    return p


def remove(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def cell_write(cell, text, header=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(text), 7.7 if header else 7.5, bold=header)


doc = Document(DOCX)

# El resumen se retira de la portada para evitar una segunda página huérfana.
front_starts = (
    'Resumen ejecutivo',
    'Aether convierte necesidades organizacionales',
    'La solución establece límites claros',
    'Esta versión deja definidos los objetivos operativos',
)
for paragraph in list(doc.paragraphs):
    if paragraph.text.startswith(front_starts):
        remove(paragraph)

# El índice general pasa a reflejar todas las secciones normativas de la SRS.
index = doc.tables[3]
for row in list(index.rows)[-1:]:
    if row.cells[0].text == 'A':
        index._tbl.remove(row._tr)
for code, detail in [
    ('8', 'Perfiles, permisos y casos de uso'),
    ('9', 'Diagramas de arquitectura, datos y estados'),
    ('10', 'Modelo de datos y contratos'),
    ('11', 'API, errores e integración'),
    ('12', 'Operación, resiliencia y recuperación'),
    ('13', 'Guía de implementación y pruebas'),
    ('14', 'Caso de negocio y aprobación formal'),
    ('A', 'Plantillas operativas y apéndices'),
]:
    cells = index.add_row().cells
    cell_write(cells[0], code)
    cell_write(cells[1], detail)
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('Tablas 1 a 2 control documental.'):
        paragraph.clear()
        set_font(paragraph.add_run('Tablas 1 a 2 control documental. Tablas 3 a 16 alcance, requisitos, datos, aceptación y trazabilidad. Tablas A1 a A4 plantillas y glosario. Tablas B1 a B14 perfiles, casos de uso, datos, operación, pruebas y caso de negocio.'), 9.3)

# Se incorpora el resumen ejecutivo al inicio del caso de negocio, antes de la decisión solicitada.
section = next(p for p in doc.paragraphs if p.text == '14 Caso de negocio y aprobación formal')
p = add_after(section, '14.1 Resumen ejecutivo', style='Heading 2', after=4)
p = add_after(p, 'Aether convierte necesidades organizacionales en iniciativas evaluables, decisiones trazables y proyectos ejecutables. La plataforma centraliza el ciclo desde la identificación de una necesidad hasta su seguimiento, evidencia y cierre, manteniendo separación explícita entre iniciativa, evaluación, decisión y ejecución.', after=4)
p = add_after(p, 'La solución establece límites claros entre navegador, aplicación web, servidor, dominio, PostgreSQL, worker, Keycloak y Redis. La seguridad se basa en OIDC con PKCE, sesiones opacas, CSRF y autorización contextual de servidor. Las operaciones relevantes se auditan y los efectos asíncronos se protegen con outbox transaccional, consumidores idempotentes y dead letters.', after=4)
add_after(p, 'Esta versión deja definidos los objetivos operativos iniciales, la política de retención, el modelo de aprobación y un caso de negocio de referencia. La autorización formal de las personas designadas sigue siendo una condición previa para construir o liberar a producción.', after=8)

renumber = {
    '14.1 Decisión solicitada': '14.2 Decisión solicitada',
    '14.2 Problema y oportunidad': '14.3 Problema y oportunidad',
    '14.3 Presupuesto de referencia V1': '14.4 Presupuesto de referencia V1',
    '14.4 Cronograma y hitos de control': '14.5 Cronograma y hitos de control',
    '14.5 Beneficio y retorno de referencia': '14.6 Beneficio y retorno de referencia',
    '14.6 Condiciones para la aprobación formal': '14.7 Condiciones para la aprobación formal',
}
for paragraph in doc.paragraphs:
    if paragraph.text in renumber:
        replacement = renumber[paragraph.text]
        paragraph.clear()
        set_font(paragraph.add_run(replacement), 9)
        paragraph.style = 'Heading 2'

# El marcador de cierre anterior se reemplaza por el cierre final de la sección 14.
ends = [p for p in doc.paragraphs if p.text == 'Fin del documento']
for paragraph in ends[:-1]:
    remove(paragraph)

doc.save(DOCX)
print(DOCX)
