from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


DOCX = Path(__file__).resolve().parent / 'Aether_SRS_2026-09.docx'


def remove(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


doc = Document(DOCX)

# La versión controlada no incluye prototipos visuales, conforme al alcance solicitado.
for paragraph in list(doc.paragraphs):
    if paragraph.text == '13.4 Alcance del diseño de interfaz' or paragraph.text.startswith('Los recorridos y criterios de interfaz se especifican'):
        remove(paragraph)

# Se evita que las condiciones de aprobación queden comprimidas después de la tabla financiera.
heading = next(p for p in doc.paragraphs if p.text == '14.7 Condiciones para la aprobación formal')
break_paragraph = OxmlElement('w:p')
break_run = OxmlElement('w:r')
page_break = OxmlElement('w:br')
page_break.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', 'page')
break_run.append(page_break)
break_paragraph.append(break_run)
heading._p.addprevious(break_paragraph)

doc.save(DOCX)
print(DOCX)
